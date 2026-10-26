from flask import Flask, request, jsonify
from flask_cors import CORS
import os
import json
import logging
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)
CORS(app)

try:
    import google.generativeai as genai
except ImportError:
    genai = None

from models.embedder import EmbedderModel

embedder = None
embedder_status = "uninitialized"
embedder_error = None
enable_local_embedder = os.getenv("ENABLE_LOCAL_EMBEDDER", "false").lower() == "true"

def init_embedder():
    global embedder, embedder_status, embedder_error
    if enable_local_embedder:
        try:
            logger.info("ENABLE_LOCAL_EMBEDDER is true. Initializing sentence-transformers embedder...")
            embedder = EmbedderModel()
            embedder_status = "loaded"
            embedder_error = None
            logger.info("Embedder model loaded successfully with backend: %s", embedder.backend)
        except Exception as e:
            embedder = None
            embedder_status = "failed"
            embedder_error = str(e)
            logger.error("Failed to load sentence-transformers model: %s", e)
    else:
        embedder = None
        embedder_status = "disabled"
        embedder_error = "ENABLE_LOCAL_EMBEDDER is set to false. Embedder is disabled."
        logger.info("ENABLE_LOCAL_EMBEDDER is false. Embedding generation is disabled.")

init_embedder()

# Load Gemini
gemini_key = os.getenv("GEMINI_API_KEY")
if gemini_key and genai:
    genai.configure(api_key=gemini_key)

def heuristic_parse(text: str) -> dict:
    import re
    emails = re.findall(r"[\w.+-]+@[\w-]+\.\w+", text)
    phones = re.findall(r"[\+\(]?[1-9][0-9 .\-\(\)]{8,}[0-9]", text)
    return {
        "contact": {
            "email": emails[0] if emails else None,
            "phone": phones[0] if phones else None,
        },
        "summary": None,
        "skills": [],
        "experience": [],
        "education": [],
        "projects": [],
        "certifications": [],
    }

@app.route("/health", methods=["GET"])
def health():
    is_healthy = embedder_status == "loaded" if enable_local_embedder else True
    return jsonify({
        "status": "ok" if is_healthy else "degraded",
        "embedder": {
            "enabled": enable_local_embedder,
            "status": embedder_status,
            "model": getattr(embedder, "model_name", "sentence-transformers/all-MiniLM-L6-v2") if embedder else None,
            "dimension": getattr(embedder, "dimension", 384) if embedder else None,
            "backend": getattr(embedder, "backend", None),
            "device": getattr(embedder, "device", None),
            "error": embedder_error
        }
    }), (200 if is_healthy else 503)

@app.route("/embed", methods=["POST"])
def embed_text():
    if embedder is None or embedder_status != "loaded":
        return jsonify({
            "error": "Embedding model is unavailable. ENABLE_LOCAL_EMBEDDER must be set to 'true' and the model must be loaded.",
            "status": embedder_status,
            "detail": embedder_error
        }), 503

    req = request.get_json(silent=True) or {}
    text = req.get("text")
    if text is None or not isinstance(text, str) or not text.strip():
        return jsonify({"error": "Field 'text' is required and must not be empty or whitespace only."}), 400

    try:
        vec = embedder.embed(text)
        return jsonify({
            "embedding": vec,
            "dim": len(vec),
            "model": embedder.model_name,
            "backend": embedder.backend
        })
    except ValueError as ve:
        return jsonify({"error": str(ve)}), 400
    except Exception as e:
        logger.error("Embedding generation failed: %s", e)
        return jsonify({"error": f"Failed to generate embedding: {str(e)}"}), 500

@app.route("/embed/batch", methods=["POST"])
def embed_batch():
    if embedder is None or embedder_status != "loaded":
        return jsonify({
            "error": "Embedding model is unavailable. ENABLE_LOCAL_EMBEDDER must be set to 'true' and the model must be loaded.",
            "status": embedder_status,
            "detail": embedder_error
        }), 503

    req = request.get_json(silent=True) or {}
    texts = req.get("texts")
    if not isinstance(texts, list) or len(texts) == 0:
        return jsonify({"error": "Field 'texts' must be a non-empty list of strings."}), 400

    try:
        vecs = embedder.embed_batch(texts)
        return jsonify({
            "embeddings": vecs,
            "dim": len(vecs[0]),
            "count": len(vecs),
            "model": embedder.model_name
        })
    except ValueError as ve:
        return jsonify({"error": str(ve)}), 400
    except Exception as e:
        logger.error("Batch embedding generation failed: %s", e)
        return jsonify({"error": f"Failed to generate batch embeddings: {str(e)}"}), 500

@app.route("/parse-resume", methods=["POST"])
def parse_resume():
    if 'file' not in request.files:
        return jsonify({"error": "No file part in request"}), 400
    file = request.files['file']
    filename = file.filename or ""
    content = file.read()

    if not content:
        return jsonify({"error": "Uploaded file is empty"}), 400

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    raw_text = ""

    if ext == "pdf":
        import pdfplumber
        import io
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                pages_text = []
                for p in pdf.pages:
                    txt = p.extract_text()
                    if txt:
                        pages_text.append(txt)
                raw_text = "\n".join(pages_text)
        except Exception as e:
            logger.warning("pdfplumber failed: %s", e)
    elif ext in ("docx", "doc"):
        import io
        from docx import Document
        try:
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            raw_text = "\n".join(paragraphs)
        except Exception as e:
            logger.warning("python-docx failed: %s", e)
    elif ext in ("txt", "md"):
        raw_text = content.decode("utf-8", errors="replace")

    raw_text = EmbedderModel.clean_text(raw_text)

    if not gemini_key:
        parsed = heuristic_parse(raw_text)
    else:
        try:
            model = genai.GenerativeModel('gemini-1.5-flash')
            prompt = f"""Parse the following resume text into a JSON object with these fields:
contact (name, email, phone, location, linkedin, github), summary, skills (list of strings),
experience (list of {{company, title, start, end, bullets}}),
education (list of {{institution, degree, field, start, end, gpa}}),
projects (list of {{name, description, tech}}),
certifications (list of {{name, issuer, date}}).
Only include fields that are present. Respond with pure JSON.

Resume:
{raw_text[:10000]}"""
            resp = model.generate_content(prompt, generation_config=genai.GenerationConfig(response_mime_type="application/json"))
            parsed = json.loads(resp.text)
        except Exception as e:
            logger.warning(f"Gemini parse failed: {e}; falling back to heuristics")
            parsed = heuristic_parse(raw_text)

    return jsonify({
        "contact": parsed.get("contact", {}),
        "summary": parsed.get("summary"),
        "skills": parsed.get("skills", []),
        "experience": parsed.get("experience", []),
        "education": parsed.get("education", []),
        "projects": parsed.get("projects", []),
        "certifications": parsed.get("certifications", []),
        "raw_text": raw_text
    })

@app.route("/enhance", methods=["POST"])
def enhance_resume():
    req = request.get_json(silent=True) or {}
    resume_json = req.get("resume_json", {})
    job_desc = req.get("job_description", "")

    if not gemini_key:
        return jsonify({
            "enhanced_resume": resume_json,
            "ats_score_before": 50,
            "ats_score_after": 50,
            "changes": []
        })

    try:
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""You are an expert resume coach.
Given the resume JSON and the job description, produce a tailored version.

STRICT RULES:
- NEVER invent employers, degrees, schools, or projects the candidate didn't have.
- NEVER add skills the candidate doesn't already have listed. List them as gaps instead.
- Only rewrite bullet points with stronger action verbs and JD-relevant keywords.
- Reorder the skills list to put most-relevant first.
- Improve the summary to align with the job.
- Return a JSON object with: enhanced_resume (same schema as input), changes (list of {{section, field, before, after}}), ats_score_before (0-100), ats_score_after (0-100), gaps (list of skill strings the candidate is missing).

Job Description:
{job_desc[:4000]}

Resume JSON:
{json.dumps(resume_json)[:6000]}

Respond with pure JSON only."""
        resp = model.generate_content(prompt, generation_config=genai.GenerationConfig(response_mime_type="application/json"))
        data = json.loads(resp.text)
        return jsonify({
            "enhanced_resume": data.get("enhanced_resume", resume_json),
            "ats_score_before": data.get("ats_score_before", 50),
            "ats_score_after": data.get("ats_score_after", 60),
            "changes": data.get("changes", [])
        })
    except Exception as e:
        logger.error(f"Enhance failed: {e}")
        return jsonify({
            "enhanced_resume": resume_json,
            "ats_score_before": 50,
            "ats_score_after": 50,
            "changes": []
        })

@app.route("/skills/extract", methods=["POST"])
def extract_skills():
    req = request.get_json(silent=True) or {}
    text = req.get("text", "")
    if gemini_key:
        try:
            model = genai.GenerativeModel('gemini-1.5-flash')
            prompt = f"Extract all technical and professional skills from this text as a JSON array of strings named 'skills'. Normalize them (e.g. 'JS' -> 'JavaScript'). Text:\n{text[:4000]}"
            resp = model.generate_content(prompt, generation_config=genai.GenerationConfig(response_mime_type="application/json"))
            data = json.loads(resp.text)
            skills = data.get("skills", [])
            return jsonify({"skills": [str(s) for s in skills]})
        except Exception as e:
            logger.warning(f"LLM skill extraction failed: {e}")

    common_skills = [
        "Python", "Java", "JavaScript", "TypeScript", "React", "Node.js",
        "Spring Boot", "SQL", "PostgreSQL", "MongoDB", "Docker", "Kubernetes",
        "AWS", "GCP", "Azure", "Git", "REST API", "GraphQL", "FastAPI",
        "Machine Learning", "TensorFlow", "PyTorch", "NLP", "Agile", "Scrum",
        "Microservices", "Kafka", "Redis", "CI/CD", "Linux"
    ]
    text_lower = text.lower()
    found = [s for s in common_skills if s.lower() in text_lower]
    return jsonify({"skills": found})

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000)
